#!/usr/bin/env python3
"""Import the tournament's Results worksheet into the static site's JSON data."""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_ROUNDS = 12
EXPECTED_PLACEMENTS = set(range(1, 12))


def normalize(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", text).strip().casefold()


def placement_number(value: Any) -> int | None:
    if isinstance(value, int) and not isinstance(value, bool):
        return value if 1 <= value <= 12 else None
    text = normalize(value)
    match = re.fullmatch(
        r"\s*(\d{1,2})\s*(?:°|º|o|º|°)?\s*(?:(?:posto|classificato|classifica))?\s*",
        text,
        flags=re.IGNORECASE,
    )
    if not match:
        return None
    number = int(match.group(1))
    return number if 1 <= number <= 12 else None


def parse_date(value: Any, row_number: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        text = value.strip()
        for date_format in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(text, date_format).date().isoformat()
            except ValueError:
                continue
    raise ValueError(f"Data non riconosciuta alla riga {row_number}: {value!r}")


def find_results_layout(sheet: Any) -> tuple[int, int, dict[int, int], int]:
    for row_number in range(1, min(sheet.max_row, 40) + 1):
        row_values = {
            column: sheet.cell(row_number, column).value
            for column in range(1, sheet.max_column + 1)
        }
        headers = {
            column: normalize(value)
            for column, value in row_values.items()
        }
        placements = {
            number: column
            for column, value in row_values.items()
            if (number := placement_number(value)) is not None
        }
        joker_columns = [
            column for column, header in headers.items()
            if header in {"jolly", "jolly player", "giocatore jolly"}
        ]
        date_columns = [
            column for column, header in headers.items()
            if header in {"data", "date", "data giornata"}
        ]
        if not date_columns:
            date_columns = [
                column for column, header in headers.items()
                if header == "giornata"
            ]
        if set(placements) == EXPECTED_PLACEMENTS and len(joker_columns) == 1 and len(date_columns) == 1:
            date_column = date_columns[0] if len(date_columns) == 1 else 1
            return row_number, date_column, placements, joker_columns[0]
    raise ValueError(
        "Non trovo nell'intestazione di 'Risultati' le colonne Data, 1°–11° e Jolly. "
        "Controlla che l'intestazione sia nella stessa riga e che i nomi delle colonne "
        "siano Data, 1°, 2°…11° e Jolly."
    )


def find_matrix_layout(
    sheet: Any,
    players: list[dict[str, str]],
) -> tuple[int, dict[int, int], int, dict[str, int]]:
    player_ids = {normalize(player["name"]): player["id"] for player in players}
    for row_number in range(1, min(sheet.max_row, 40) + 1):
        if normalize(sheet.cell(row_number, 1).value) not in {"giocatore", "giocatori", "player"}:
            continue
        dates: dict[int, int] = {}
        for column in range(2, sheet.max_column + 1):
            try:
                date_value = parse_date(sheet.cell(row_number, column).value, row_number)
            except ValueError:
                continue
            dates[column] = column
        if len(dates) != EXPECTED_ROUNDS or len({parse_date(sheet.cell(row_number, column).value, row_number) for column in dates}) != EXPECTED_ROUNDS:
            continue

        player_rows: dict[str, int] = {}
        joker_rows: list[int] = []
        for candidate_row in range(row_number + 1, sheet.max_row + 1):
            label = normalize(sheet.cell(candidate_row, 1).value)
            if label.startswith("jolly") or label == "giocatore jolly":
                joker_rows.append(candidate_row)
            elif label in player_ids:
                player_rows[player_ids[label]] = candidate_row
            elif label.startswith("serata valida"):
                continue
            elif label:
                raise ValueError(f"Giocatore non riconosciuto nella matrice 'Risultati': {sheet.cell(candidate_row, 1).value!r}.")
        if len(player_rows) != len(players) or len(joker_rows) != 1:
            raise ValueError(
                "Nel foglio 'Risultati' a matrice devono comparire una volta tutti i giocatori "
                "iscritti e una sola riga 'Jolly'."
            )
        return row_number, dates, joker_rows[0], player_rows
    raise ValueError(
        "Il foglio 'Risultati' non contiene né l'intestazione classica né la matrice "
        "con una riga 'Giocatore', 12 date in colonna e una riga 'Jolly'."
    )


def matrix_results(
    sheet: Any,
    players: list[dict[str, str]],
    header_row: int,
    date_columns: dict[int, int],
    joker_row: int,
    player_rows: dict[str, int],
) -> tuple[list[str], list[dict[str, Any]]]:
    player_ids = {normalize(player["name"]): player["id"] for player in players}
    date_to_column = {
        parse_date(sheet.cell(header_row, column).value, header_row): column
        for column in date_columns
    }
    if len(date_to_column) != EXPECTED_ROUNDS:
        raise ValueError("Le 12 date della matrice devono essere tutte distinte e valide.")

    rounds: list[dict[str, Any]] = []
    joker_users: set[str] = set()
    for round_date, column in sorted(date_to_column.items()):
        placements: list[str | None] = [None] * 12
        for player_id, row_number in player_rows.items():
            value = sheet.cell(row_number, column).value
            if value in (None, ""):
                continue
            placement = placement_number(value)
            if placement is None:
                raise ValueError(
                    f"Piazzamento non valido per {sheet.cell(row_number, 1).value!r} "
                    f"nella giornata del {round_date}: {value!r}"
                )
            if placements[placement - 1] is not None:
                raise ValueError(
                    f"Piazzamento {placement} duplicato nella giornata del {round_date}."
                )
            placements[placement - 1] = player_id

        joker_value = sheet.cell(joker_row, column).value
        joker_player_id = None
        if joker_value not in (None, ""):
            joker_player_id = player_ids.get(normalize(joker_value))
            if joker_player_id is None:
                raise ValueError(
                    f"Giocatore jolly non riconosciuto nella giornata del {round_date}: {joker_value!r}"
                )
            if joker_player_id not in placements:
                raise ValueError(
                    f"Il giocatore jolly {joker_value!r} non ha un piazzamento "
                    f"nella giornata del {round_date}."
                )
            if joker_player_id in joker_users:
                raise ValueError(f"{joker_value!r} risulta aver usato il jolly più di una volta.")
            joker_users.add(joker_player_id)

        rounds.append({
            "date": round_date,
            "placements": placements,
            "jokerPlayerId": joker_player_id,
        })

    return list(sorted(date_to_column)), rounds


def workbook_results(workbook_path: Path, players: list[dict[str, str]]) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet_names = {normalize(sheet_name): sheet_name for sheet_name in workbook.sheetnames}
    if "risultati" not in sheet_names:
        raise ValueError("Il file Excel deve contenere un foglio chiamato 'Risultati'.")

    sheet = workbook[sheet_names["risultati"]]
    matrix_layout: tuple[int, dict[int, int], int, dict[str, int]] | None = None
    try:
        matrix_layout = find_matrix_layout(sheet, players)
    except ValueError as matrix_error:
        try:
            find_results_layout(sheet)
        except ValueError:
            raise matrix_error
    if matrix_layout is not None:
        try:
            return matrix_results(sheet, players, *matrix_layout)
        finally:
            workbook.close()

    header_row, date_column, placement_columns, joker_column = find_results_layout(sheet)
    player_ids = {normalize(player["name"]): player["id"] for player in players}
    rows_by_date: dict[str, dict[str, Any]] = {}
    joker_users: set[str] = set()

    for row_number in range(header_row + 1, sheet.max_row + 1):
        row = {
            column: sheet.cell(row_number, column).value
            for column in range(1, sheet.max_column + 1)
        }
        date_value = row.get(date_column)
        has_values = any(value not in (None, "") for value in row.values())
        if date_value in (None, ""):
            if has_values:
                raise ValueError(
                    f"La riga {row_number} di 'Risultati' contiene dati ma non ha una data."
                )
            continue

        round_date = parse_date(date_value, row_number)
        if round_date in rows_by_date:
            raise ValueError(f"La data {round_date} è ripetuta nel foglio 'Risultati'.")

        placements: list[str | None] = [None] * 12
        for placement, column in placement_columns.items():
            value = row.get(column)
            if value in (None, ""):
                continue
            player_id = player_ids.get(normalize(value))
            if player_id is None:
                raise ValueError(
                    f"Nome giocatore sconosciuto alla riga {row_number}, "
                    f"piazzamento {placement}: {value!r}"
                )
            placements[placement - 1] = player_id

        filled = [player_id for player_id in placements if player_id is not None]
        if len(set(filled)) != len(filled):
            raise ValueError(f"Un giocatore è ripetuto nella serata del {round_date}.")

        joker_value = row.get(joker_column)
        joker_player_id = None
        if joker_value not in (None, ""):
            joker_player_id = player_ids.get(normalize(joker_value))
            if joker_player_id is None:
                raise ValueError(
                    f"Il valore della colonna Jolly alla riga {row_number} "
                    f"non corrisponde a un giocatore: {joker_value!r}"
                )
            if joker_player_id not in filled:
                raise ValueError(
                    f"Il giocatore con il jolly non ha un piazzamento nella serata del {round_date}."
                )
            if joker_player_id in joker_users:
                raise ValueError(f"{joker_value!r} risulta aver usato il jolly più di una volta.")
            joker_users.add(joker_player_id)

        rows_by_date[round_date] = {
            "date": round_date,
            "placements": placements,
            "jokerPlayerId": joker_player_id,
        }

    workbook.close()
    if len(rows_by_date) != EXPECTED_ROUNDS:
        raise ValueError(
            f"Il foglio 'Risultati' deve contenere tutte le {EXPECTED_ROUNDS} date "
            f"del torneo; trovate {len(rows_by_date)}."
        )
    dates = sorted(rows_by_date)
    return dates, [rows_by_date[round_date] for round_date in dates]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True, help="Percorso al file Excel.")
    parser.add_argument("--data", type=Path, required=True, help="Percorso a data/torneo.json.")
    args = parser.parse_args()

    try:
        data = json.loads(args.data.read_text(encoding="utf-8"))
        dates, rounds = workbook_results(args.workbook, data["players"])
        data["calendar"] = dates
        data["rounds"] = rounds
        args.data.write_text(
            json.dumps(data, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    except (OSError, KeyError, ValueError) as error:
        print(f"Errore importando il foglio Excel: {error}", file=sys.stderr)
        return 1

    print(f"Importate {len(rounds)} giornate da {args.workbook}.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
