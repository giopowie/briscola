import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from excel_to_json import workbook_results


PLAYERS = [
    {"id": "fanfo", "name": "Fanfo"},
    {"id": "generale", "name": "Generale"},
    {"id": "giopowie", "name": "Giopowie"},
    {"id": "lorenzo", "name": "Lorenzo"},
    {"id": "maldo", "name": "Maldo"},
    {"id": "muso", "name": "Muso"},
    {"id": "piero", "name": "Piero"},
    {"id": "salado", "name": "Salado"},
    {"id": "torre", "name": "Torre"},
    {"id": "ube", "name": "Ube"},
    {"id": "uccio", "name": "Uccio"},
]
DATES = [
    "2026-10-29",
    "2026-11-12",
    "2026-11-26",
    "2026-12-10",
    "2027-01-07",
    "2027-01-21",
    "2027-02-04",
    "2027-02-18",
    "2027-03-04",
    "2027-03-18",
    "2027-04-01",
    "2027-04-15",
]


def save_workbook(workbook):
    temporary_file = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    temporary_file.close()
    workbook.save(temporary_file.name)
    return Path(temporary_file.name)


class ExcelImporterTests(unittest.TestCase):
    def make_matrix(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Risultati"
        sheet.cell(1, 1, "Giocatore")
        for column, round_date in enumerate(DATES, start=2):
            year, month, day = map(int, round_date.split("-"))
            sheet.cell(1, column, date(year, month, day))
        for row, player in enumerate(PLAYERS, start=2):
            sheet.cell(row, 1, player["name"])
        sheet.cell(len(PLAYERS) + 2, 1, "Jolly")
        sheet.cell(len(PLAYERS) + 3, 1, "Serata valida (5+)")
        return workbook, sheet

    def test_imports_player_rows_and_dates_as_columns(self):
        workbook, sheet = self.make_matrix()
        for placement in range(1, 6):
            sheet.cell(placement + 1, 2, placement)
        sheet.cell(len(PLAYERS) + 2, 2, "Fanfo")
        workbook_path = save_workbook(workbook)
        try:
            dates, rounds = workbook_results(workbook_path, PLAYERS)
        finally:
            workbook_path.unlink()

        self.assertEqual(len(dates), 12)
        self.assertEqual(rounds[0]["date"], "2026-10-29")
        self.assertEqual(rounds[0]["placements"][:5], [player["id"] for player in PLAYERS[:5]])
        self.assertEqual(rounds[0]["jokerPlayerId"], "fanfo")
        self.assertTrue(all(not any(round_item["placements"]) for round_item in rounds[1:]))

    def test_rejects_duplicate_placements_in_matrix(self):
        workbook, sheet = self.make_matrix()
        sheet.cell(2, 2, 1)
        sheet.cell(3, 2, 1)
        workbook_path = save_workbook(workbook)
        try:
            with self.assertRaisesRegex(ValueError, "Piazzamento 1 duplicato"):
                workbook_results(workbook_path, PLAYERS)
        finally:
            workbook_path.unlink()

    def test_keeps_support_for_original_results_layout(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Risultati"
        headers = ["Giornata", "Data", *(f"{position}°" for position in range(1, 12)), "Jolly"]
        for column, header in enumerate(headers, start=1):
            sheet.cell(1, column, header)
        for row, round_date in enumerate(DATES, start=2):
            sheet.cell(row, 1, row - 1)
            year, month, day = map(int, round_date.split("-"))
            sheet.cell(row, 2, date(year, month, day))
        workbook_path = save_workbook(workbook)
        try:
            dates, rounds = workbook_results(workbook_path, PLAYERS)
        finally:
            workbook_path.unlink()

        self.assertEqual(len(dates), 12)
        self.assertTrue(all(not any(round_item["placements"]) for round_item in rounds))


if __name__ == "__main__":
    unittest.main()
